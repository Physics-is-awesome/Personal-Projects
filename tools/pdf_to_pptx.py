#!/usr/bin/env python3

from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.dml.color import RGBColor

path = input("Insert path here")
# Convert PDF pages to images
pages = convert_from_path(path, dpi=200)

# Create PowerPoint
prs = Presentation()
prs.slide_width  = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]  # blank

for page in pages:
    slide = prs.slides.add_slide(blank_layout)
    # Save temporary image
    page.save("temp_slide.png", "PNG")
    # Add full-slide image
    slide.shapes.add_picture(
        "temp_slide.png",
        left=0, top=0,
        width=prs.slide_width,
        height=prs.slide_height
    )

prs.save("velocity.pptx")
print("Done → velocity.pptx")
